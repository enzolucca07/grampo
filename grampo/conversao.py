"""Conversões: Office, imagens e texto para PDF; PDF para JPG e para Word.

Word, Excel e PowerPoint usam o Microsoft Office instalado (pelo COM do Windows); sem ele, o LibreOffice. Sem
nenhum dos dois, a tela explica o que instalar. Nada vai para a internet.
"""

import html
import os
import re
import shutil
import subprocess
import tempfile
import threading
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import pymupdf
from PIL import Image, ImageOps

from .erros import ErroAmigavel
from .operacoes import _nada, abrir_pdf

EXTENSOES_WORD = {".doc", ".docx", ".rtf", ".odt"}
EXTENSOES_EXCEL = {".xls", ".xlsx", ".xlsm", ".ods", ".csv"}
EXTENSOES_POWERPOINT = {".ppt", ".pptx", ".pps", ".ppsx", ".odp"}
EXTENSOES_IMAGEM = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
EXTENSOES_TEXTO = {".txt"}

PROGRAMA_DO_OFFICE = {"word": "Word.Application", "excel": "Excel.Application", "powerpoint": "PowerPoint.Application"}
NOME_DO_PROGRAMA = {"word": "Word", "excel": "Excel", "powerpoint": "PowerPoint"}


def tipo_office(caminho):
    ext = Path(caminho).suffix.lower()
    for tipo, extensoes in (("word", EXTENSOES_WORD), ("excel", EXTENSOES_EXCEL),
                            ("powerpoint", EXTENSOES_POWERPOINT)):
        if ext in extensoes:
            return tipo
    return None


# ---------------------------------------------------------------- o que está instalado


def office_instalado(tipo):
    """O Windows conhece o Word/Excel/PowerPoint (registro do COM)?"""
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, PROGRAMA_DO_OFFICE[tipo] + r"\CLSID"):
            return True
    except (ImportError, OSError):
        return False


def achar_libreoffice():
    candidatos = [shutil.which("soffice")]
    for variavel in ("ProgramFiles", "ProgramFiles(x86)"):
        if os.environ.get(variavel):
            candidatos.append(Path(os.environ[variavel]) / "LibreOffice" / "program" / "soffice.exe")
    for caminho in candidatos:
        if caminho and Path(caminho).is_file():
            return Path(caminho)
    return None


def programas_disponiveis():
    """Para a tela: com o que cada tipo de arquivo do Office vai ser convertido (ou None)."""
    libre = achar_libreoffice()
    return {tipo: ("Microsoft " + NOME_DO_PROGRAMA[tipo]) if office_instalado(tipo) else ("LibreOffice" if libre else None)
            for tipo in PROGRAMA_DO_OFFICE}


# ---------------------------------------------------------------- Office -> PDF


IMPRESSORAS_DE_PDF = ("Microsoft Print to PDF", "Microsoft XPS Document Writer")
_TRAVA_DA_IMPRESSORA = threading.Lock()


def _resolucao(impressora):
    """Pontos por polegada na horizontal. Impressora "Somente Texto" (como a do Domínio/Domsis) dá uns 10."""
    import win32ui

    contexto = win32ui.CreateDC()
    try:
        contexto.CreatePrinterDC(impressora)
        return contexto.GetDeviceCaps(88)  # LOGPIXELSX
    finally:
        contexto.DeleteDC()


def devolver_impressora():
    """Se o programa fechou no meio de uma conversão, a impressora padrão ainda está trocada: devolve."""
    from .saida import gravar_configuracao, ler_configuracao

    dados = ler_configuracao()
    original = dados.pop("impressora_para_devolver", None)
    if original:
        try:
            import win32print
            win32print.SetDefaultPrinter(original)
        except Exception:
            return
        gravar_configuracao(dados)


@contextmanager
def impressora_que_desenha():
    """O Word e o Excel montam as páginas com as medidas da impressora padrão. Se ela é "Somente Texto", a letra
    sai enorme no PDF; então, só durante a conversão, a padrão vira a "Microsoft Print to PDF" (e depois volta)."""
    from .saida import gravar_configuracao, ler_configuracao

    with _TRAVA_DA_IMPRESSORA:
        troca = original = None
        try:
            import win32print

            original = win32print.GetDefaultPrinter()
            if _resolucao(original) < 150:
                instaladas = {p[2] for p in win32print.EnumPrinters(win32print.PRINTER_ENUM_LOCAL)}
                troca = next((p for p in IMPRESSORAS_DE_PDF if p in instaladas), None)
        except Exception:
            troca = None
        if not troca:
            yield
            return
        gravar_configuracao({**ler_configuracao(), "impressora_para_devolver": original})
        win32print.SetDefaultPrinter(troca)
        try:
            yield
        finally:
            win32print.SetDefaultPrinter(original)
            dados = ler_configuracao()
            dados.pop("impressora_para_devolver", None)
            gravar_configuracao(dados)


def _pelo_office(tipo, entrada, destino):
    with impressora_que_desenha():
        return _pelo_office_agora(tipo, entrada, destino)


def _pelo_office_agora(tipo, entrada, destino):
    import pythoncom
    import win32com.client

    entrada, destino = str(Path(entrada).resolve()), str(Path(destino).resolve())
    pythoncom.CoInitialize()
    programa = None
    try:
        programa = win32com.client.DispatchEx(PROGRAMA_DO_OFFICE[tipo])
        if tipo == "word":
            programa.Visible = False
            programa.DisplayAlerts = 0
            doc = programa.Documents.Open(entrada, ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False,
                                          PasswordDocument="\x00", Visible=False)
            try:
                doc.ExportAsFixedFormat(destino, 17)  # 17 = PDF
            finally:
                doc.Close(0)
        elif tipo == "excel":
            programa.Visible = False
            programa.DisplayAlerts = False
            pasta = programa.Workbooks.Open(entrada, UpdateLinks=0, ReadOnly=True, Password="\x00",
                                            AddToMru=False)
            try:
                pasta.ExportAsFixedFormat(0, destino)  # 0 = PDF, todas as planilhas com conteúdo
            finally:
                pasta.Close(False)
        else:
            apresentacao = programa.Presentations.Open(entrada, ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                apresentacao.SaveAs(destino, 32)  # 32 = PDF
            finally:
                apresentacao.Close()
    except Exception as erro:
        texto = str(erro).lower()
        if "senha" in texto or "password" in texto:
            raise ErroAmigavel(f"“{Path(entrada).name}” tem senha. Abra no {NOME_DO_PROGRAMA[tipo]}, tire a senha "
                               "e tente de novo.") from None
        raise ErroAmigavel(f"O {NOME_DO_PROGRAMA[tipo]} não conseguiu converter “{Path(entrada).name}”. Abra o "
                           "arquivo nele para ver se está tudo certo e tente de novo.") from None
    finally:
        if programa is not None:
            try:
                programa.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
    return Path(destino)


def _pelo_libreoffice(soffice, entrada, destino):
    with tempfile.TemporaryDirectory(prefix="Grampo-lo-") as temporaria:
        perfil = Path(temporaria) / "perfil"  # perfil próprio: funciona mesmo com o LibreOffice já aberto
        comando = [str(soffice), f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--convert-to", "pdf",
                   "--outdir", temporaria, str(entrada)]
        try:
            subprocess.run(comando, capture_output=True, timeout=300,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except subprocess.TimeoutExpired:
            raise ErroAmigavel(f"O LibreOffice demorou demais para converter “{Path(entrada).name}”.") from None
        gerado = Path(temporaria) / (Path(entrada).stem + ".pdf")
        if not gerado.is_file():
            raise ErroAmigavel(f"O LibreOffice não conseguiu converter “{Path(entrada).name}”.")
        Path(destino).parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(gerado), destino)
    return Path(destino)


def office_para_pdf(entrada, destino):
    tipo = tipo_office(entrada)
    if tipo is None:
        raise ErroAmigavel(f"“{Path(entrada).name}” não é um arquivo do Word, Excel ou PowerPoint.")
    if office_instalado(tipo):
        return _pelo_office(tipo, entrada, destino)
    soffice = achar_libreoffice()
    if soffice:
        return _pelo_libreoffice(soffice, entrada, destino)
    raise ErroAmigavel(f"Para converter arquivos do {NOME_DO_PROGRAMA[tipo]} é preciso ter o Microsoft Office ou o "
                       "LibreOffice (gratuito) instalado neste computador. Peça para a TI instalar um deles.")


# ---------------------------------------------------------------- imagens e texto -> PDF

A4 = pymupdf.paper_rect("a4")


def _imagem_pronta(caminho):
    """Abre a imagem já em pé (foto de celular vem deitada no arquivo) e sem transparência."""
    try:
        with Image.open(caminho) as original:
            imagem = ImageOps.exif_transpose(original)
            imagem.load()
    except Exception:
        raise ErroAmigavel(f"Não consegui abrir “{Path(caminho).name}” como imagem.") from None
    if imagem.mode in ("RGBA", "LA", "P"):
        fundo = Image.new("RGB", imagem.size, "white")
        imagem = imagem.convert("RGBA")
        fundo.paste(imagem, mask=imagem.getchannel("A"))
        imagem = fundo
    elif imagem.mode not in ("RGB", "L"):
        imagem = imagem.convert("RGB")
    return imagem


def imagens_para_pdf(entradas, destino, tamanho="a4", margem=True, avisar=_nada):
    """Cada imagem vira uma página. `tamanho`: "a4" (deitada ou em pé conforme a imagem) ou "imagem"."""
    if not entradas:
        raise ErroAmigavel("Escolha pelo menos uma imagem.")
    with pymupdf.open() as doc:
        for i, entrada in enumerate(entradas):
            avisar(i / len(entradas), f"Imagem {i + 1} de {len(entradas)}")
            imagem = _imagem_pronta(entrada)
            largura, altura = imagem.size
            temporario = tempfile.SpooledTemporaryFile()
            formato = "PNG" if Path(entrada).suffix.lower() == ".png" else "JPEG"
            imagem.save(temporario, formato, **({"quality": 92} if formato == "JPEG" else {}))
            temporario.seek(0)
            if tamanho == "imagem":
                dpi = (imagem.info.get("dpi") or (96, 96))[0] or 96
                folha = doc.new_page(width=largura * 72 / dpi, height=altura * 72 / dpi)
                espaco = folha.rect
            else:
                deitada = largura > altura
                folha = doc.new_page(width=A4.height if deitada else A4.width,
                                     height=A4.width if deitada else A4.height)
                espaco = folha.rect + ((28, 28, -28, -28) if margem else (0, 0, 0, 0))
            folha.insert_image(espaco, stream=temporario.read(), keep_proportion=True)
        avisar(0.95, "Gravando")
        doc.save(destino, garbage=3, deflate=True)
    return Path(destino)


def _ler_texto(caminho):
    dados = Path(caminho).read_bytes()
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1")


def texto_para_pdf(entrada, destino):
    texto = _ler_texto(entrada)
    corpo = (f'<pre style="font-family: monospace; font-size: 10pt; white-space: pre-wrap;">'
             f"{html.escape(texto) or ' '}</pre>")
    historia = pymupdf.Story(html=corpo)
    escritor = pymupdf.DocumentWriter(str(destino))
    area = A4 + (50, 50, -50, -50)
    mais = True
    while mais:
        dispositivo = escritor.begin_page(A4)
        mais, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    return Path(destino)


# ---------------------------------------------------------------- PDF -> imagem, Word, Excel, PowerPoint, texto


def pdf_para_jpg(entrada, pasta, nome_base, dpi=150, avisar=_nada, formato="jpg"):
    """Cada página vira uma imagem (`formato`: "jpg" ou "png")."""
    formato = "png" if formato == "png" else "jpg"
    saidas = []
    with abrir_pdf(entrada) as doc:
        digitos = len(str(doc.page_count))
        for i, folha in enumerate(doc):
            avisar(i / doc.page_count, f"Página {i + 1} de {doc.page_count}")
            imagem = folha.get_pixmap(dpi=int(dpi), alpha=False)
            nome = nome_base if doc.page_count == 1 else f"{nome_base} - página {i + 1:0{digitos}d}"
            destino = Path(pasta) / f"{nome}.{formato}"
            destino.parent.mkdir(parents=True, exist_ok=True)
            imagem.save(destino, **({"jpg_quality": 90} if formato == "jpg" else {}))
            saidas.append(destino)
    return saidas


def _sem_texto(entrada):
    return ErroAmigavel(f"“{Path(entrada).name}” parece escaneado (é uma foto do papel, sem texto). Não dá para "
                        "tirar o texto dele sem reconhecimento de letras (OCR).")


def pdf_para_texto(entrada, destino, avisar=_nada):
    partes, algum = [], False
    with abrir_pdf(entrada) as doc:
        for i, folha in enumerate(doc):
            avisar(i / doc.page_count, f"Página {i + 1} de {doc.page_count}")
            texto = folha.get_text(sort=True).strip()
            algum |= bool(texto)
            partes.append(f"--- Página {i + 1} ---\n{texto}" if doc.page_count > 1 else texto)
    if not algum:
        raise _sem_texto(entrada)
    # utf-8 com BOM: o Bloco de Notas antigo mostra os acentos certo
    Path(destino).write_text("\n\n".join(partes) + "\n", encoding="utf-8-sig")
    return Path(destino)


NUMERO_BR = re.compile(r"^\(?-?\s*(R\$)?\s*-?\d{1,3}(\.\d{3})*(,\d+)?\)?%?$|^\(?-?\s*(R\$)?\s*-?\d+(,\d+)?\)?%?$")


def _celula(texto):
    """"1.234,56" vira o número 1234.56 (a planilha soma); o resto fica como texto."""
    if texto is None:
        return None, None
    texto = " ".join(str(texto).split())
    if not texto or not NUMERO_BR.match(texto):
        return texto, None
    negativo = texto.startswith("(") and texto.endswith(")") or "-" in texto
    porcento = texto.endswith("%")
    limpo = re.sub(r"[^\d,]", "", texto).replace(",", ".")
    try:
        valor = float(limpo)
    except ValueError:
        return texto, None
    if negativo:
        valor = -valor
    if porcento:
        return valor / 100, "0.00%"
    if "," in texto:
        return valor, "#,##0.00"
    if len(limpo) > 1 and limpo.startswith("0"):
        return texto, None  # código com zero na frente (CEP, conta): continua texto
    return int(valor) if valor.is_integer() and abs(valor) < 1e15 else valor, "#,##0" if "." in texto else None


def _tabelas_desenhadas(folha):
    """Tabelas com as linhas da grade desenhadas no PDF. Se as "células" vêm com várias palavras em linhas
    diferentes (colunas grudadas, comum em PDF do Excel sem grade), não confia e devolve []."""
    if not folha.get_drawings():
        return []
    try:
        # o texto de cada célula vem da área dela na página: o extract() embaralha "1.200,50"
        tabelas = [[[folha.get_textbox(pymupdf.Rect(c)).strip() if c else None for c in linha.cells]
                    for linha in t.rows] for t in folha.find_tables().tables]
    except Exception:
        return []
    boas = []
    for linhas in tabelas:
        celulas = [c for linha in linhas for c in linha if c]
        if celulas and sum("\n" in c for c in celulas) <= len(celulas) / 4:
            boas.append(_sem_vazios([[c.replace("\n", " ") if c else None for c in linha] for linha in linhas]))
    return boas


def _sem_vazios(linhas):
    """Tira colunas e linhas sem nada (aparecem quando a grade do PDF tem traços duplicados)."""
    linhas = [linha for linha in linhas if any(linha)]
    largura = max((len(linha) for linha in linhas), default=0)
    linhas = [linha + [None] * (largura - len(linha)) for linha in linhas]
    usadas = [k for k in range(largura) if any(linha[k] for linha in linhas)]
    return [[linha[k] for k in usadas] for linha in linhas]


def _por_alinhamento(folha, espaco=6, mesma_linha=3):
    """Monta linhas e colunas pela posição das palavras: espaço grande entre palavras separa células, e células
    que se sobrepõem na horizontal, em linhas diferentes, são da mesma coluna. Devolve (linhas, nº de colunas)."""
    palavras = sorted(folha.get_text("words", sort=True), key=lambda p: ((p[1] + p[3]) / 2, p[0]))
    linhas_de_palavras = []
    for p in palavras:
        meio = (p[1] + p[3]) / 2
        if linhas_de_palavras and abs(linhas_de_palavras[-1][0] - meio) <= mesma_linha:
            linhas_de_palavras[-1][1].append(p)
        else:
            linhas_de_palavras.append([meio, [p]])
    linhas = []  # cada linha: [(x0, x1, texto), ...]
    for _, grupo in linhas_de_palavras:
        grupo.sort(key=lambda p: p[0])
        altura = sorted(p[3] - p[1] for p in grupo)[len(grupo) // 2]
        limite = max(espaco, altura * 0.6)  # letra grande tem espaço grande entre as palavras
        celulas = [[grupo[0][0], grupo[0][2], grupo[0][4]]]
        for p in grupo[1:]:
            if p[0] - celulas[-1][1] > limite:
                celulas.append([p[0], p[2], p[4]])
            else:
                celulas[-1][1], celulas[-1][2] = p[2], celulas[-1][2] + " " + p[4]
        linhas.append(celulas)
    # colunas: junta as faixas horizontais (só das linhas com mais de uma célula; título comprido não conta)
    if sum(len(linha) > 1 for linha in linhas) < 2:  # tabela precisa de pelo menos duas linhas com colunas
        return [[" ".join(c[2] for c in linha)] for linha in linhas], 1
    faixas = sorted((c[0], c[1]) for linha in linhas if len(linha) > 1 for c in linha)
    colunas = []
    for x0, x1 in faixas:
        if colunas and x0 <= colunas[-1][1]:
            colunas[-1][1] = max(colunas[-1][1], x1)
        else:
            colunas.append([x0, x1])
    if len(colunas) < 2:
        return [[c[2] for c in linha] for linha in linhas], 1

    def coluna_de(x0, x1):
        sobra = [min(x1, b) - max(x0, a) for a, b in colunas]
        melhor = max(range(len(colunas)), key=lambda k: sobra[k])
        if sobra[melhor] > 0:
            return melhor
        return min(range(len(colunas)), key=lambda k: abs(colunas[k][0] - x0))

    resultado = []
    for linha in linhas:
        valores = [None] * len(colunas)
        for x0, x1, texto in linha:
            k = coluna_de(x0, x1)
            valores[k] = texto if valores[k] is None else valores[k] + " " + texto
        resultado.append(valores)
    return resultado, len(colunas)


def pdf_para_excel(entrada, destino, avisar=_nada):
    """As tabelas de cada página viram abas da planilha, com os números já como números.

    Usa a grade desenhada quando há; senão, as colunas pelo alinhamento do texto. Página só com texto corrido
    entra como texto, uma linha por linha. Devolve (destino, quantas tabelas achou)."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    planilha = Workbook()
    planilha.remove(planilha.active)
    tabelas = 0
    tem_algum_texto = False
    with abrir_pdf(entrada) as doc:
        for i, folha in enumerate(doc):
            avisar(i / doc.page_count, f"Procurando tabelas na página {i + 1} de {doc.page_count}")
            if not folha.get_text().strip():
                continue
            tem_algum_texto = True
            desenhadas = _tabelas_desenhadas(folha)
            for k, linhas in enumerate(desenhadas, start=1):
                tabelas += 1
                nome = f"Pág {i + 1}" + (f" - tabela {k}" if len(desenhadas) > 1 else "")
                _escrever_aba(planilha.create_sheet(nome[:31]), linhas, cabecalho=True)
            if desenhadas:
                continue
            linhas, colunas = _por_alinhamento(folha)
            if colunas > 1:
                tabelas += 1
                primeira = linhas[0]
                cabecalho = all(v and _celula(v)[1] is None and not NUMERO_BR.match(v) for v in primeira)
                _escrever_aba(planilha.create_sheet(f"Pág {i + 1}"), linhas, cabecalho=cabecalho)
            else:
                _escrever_aba(planilha.create_sheet(f"Pág {i + 1} (texto)"), linhas, cabecalho=False)
    if not tem_algum_texto:
        raise _sem_texto(entrada)
    for aba in planilha.worksheets:
        for coluna in aba.columns:
            maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            aba.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(10, maior + 2), 60)
    avisar(0.95, "Gravando a planilha")
    planilha.save(destino)
    return Path(destino), tabelas


def _escrever_aba(aba, linhas, cabecalho):
    """`cabecalho`: a primeira linha é o título das colunas (negrito, fica parada ao rolar)."""
    from openpyxl.styles import Font

    for numero, linha in enumerate(linhas, start=1):
        for coluna, bruto in enumerate(linha, start=1):
            valor, formato = _celula(bruto)
            celula = aba.cell(row=numero, column=coluna, value=valor)
            if formato:
                celula.number_format = formato
            if cabecalho and numero == 1:
                celula.font = Font(bold=True)
    if cabecalho:
        aba.freeze_panes = "A2"


def pdf_para_powerpoint(entrada, destino, dpi=150, avisar=_nada):
    """Cada página vira um slide (como imagem, do jeito que está no PDF)."""
    from pptx import Presentation
    from pptx.util import Emu

    apresentacao = Presentation()
    em_branco = apresentacao.slide_layouts[6]
    with abrir_pdf(entrada) as doc:
        primeira = doc[0].rect
        # slide no formato da primeira página: 25,4 cm no lado maior
        lado_maior = Emu(9144000)
        if primeira.width >= primeira.height:
            apresentacao.slide_width, apresentacao.slide_height = lado_maior, int(lado_maior * primeira.height / primeira.width)
        else:
            apresentacao.slide_height, apresentacao.slide_width = lado_maior, int(lado_maior * primeira.width / primeira.height)
        largura, altura = apresentacao.slide_width, apresentacao.slide_height
        for i, folha in enumerate(doc):
            avisar(i / doc.page_count, f"Página {i + 1} de {doc.page_count}")
            png = folha.get_pixmap(dpi=int(dpi), alpha=False).tobytes("png")
            escala = min(largura / folha.rect.width, altura / folha.rect.height)
            w, h = int(folha.rect.width * escala), int(folha.rect.height * escala)
            slide = apresentacao.slides.add_slide(em_branco)
            slide.shapes.add_picture(BytesIO(png), (largura - w) // 2, (altura - h) // 2, w, h)
    avisar(0.95, "Gravando a apresentação")
    apresentacao.save(destino)
    return Path(destino)


def extrair_imagens(entrada, pasta, nome_base, avisar=_nada):
    """As fotos e figuras que estão dentro do PDF, no tamanho original. Ignora enfeites minúsculos."""
    saidas, vistas = [], set()
    with abrir_pdf(entrada) as doc:
        for i, folha in enumerate(doc):
            avisar(i / doc.page_count, f"Página {i + 1} de {doc.page_count}")
            for info in folha.get_images(full=True):
                xref = info[0]
                if xref in vistas:
                    continue
                vistas.add(xref)
                try:
                    imagem = doc.extract_image(xref)
                except Exception:
                    continue
                if not imagem or min(imagem.get("width", 0), imagem.get("height", 0)) < 40:
                    continue
                extensao = imagem.get("ext", "png")
                if extensao not in ("jpg", "jpeg", "png"):  # formatos que nem todo programa abre: vira PNG
                    pixmap = pymupdf.Pixmap(doc, xref)
                    if pixmap.n - pixmap.alpha > 3:
                        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pixmap)
                    dados, extensao = pixmap.tobytes("png"), "png"
                else:
                    dados = imagem["image"]
                destino = Path(pasta) / f"{nome_base} - imagem {len(saidas) + 1:02d} (página {i + 1}).{extensao}"
                destino.parent.mkdir(parents=True, exist_ok=True)
                destino.write_bytes(dados)
                saidas.append(destino)
    if not saidas:
        raise ErroAmigavel(f"“{Path(entrada).name}” não tem fotos ou figuras dentro. Para ter cada página como "
                           "imagem, use “PDF para imagem”.")
    return saidas


def tem_texto(entrada):
    with abrir_pdf(entrada) as doc:
        return any(folha.get_text().strip() for folha in doc)


def pdf_para_word(entrada, destino, avisar=_nada):
    """Usa o pdf2docx (sem internet). PDF escaneado não tem texto: vira um Word com as páginas como imagem."""
    import logging

    from pdf2docx import Converter

    abrir_pdf(entrada).close()  # confere senha e se é PDF mesmo, com a mensagem certa
    logging.getLogger().setLevel(logging.ERROR)  # o pdf2docx escreve muito no console
    avisar(0.1, "Lendo o texto e o layout das páginas")
    conversor = Converter(str(entrada))
    try:
        conversor.convert(str(destino))
    except Exception:
        raise ErroAmigavel(f"Não consegui converter “{Path(entrada).name}” para Word. PDFs com layout muito "
                           "complicado às vezes não dão certo.") from None
    finally:
        conversor.close()
    return Path(destino)
