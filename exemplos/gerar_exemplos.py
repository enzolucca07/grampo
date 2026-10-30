"""Cria arquivos de exemplo para testar cada ferramenta (PDFs, imagens, texto, Word, Excel e PowerPoint).

  .venv\\Scripts\\python.exe exemplos\\gerar_exemplos.py [pasta]

Sem pasta, grava em exemplos\\arquivos. Os testes automáticos usam as mesmas funções.
"""

import sys
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw

PASTA_PADRAO = Path(__file__).resolve().parent / "arquivos"


def pdf_com_paginas(destino, paginas=5, titulo="Relatório", com_imagem=False, deitada_em=()):
    """PDF com texto "Página N" em cada folha; `deitada_em` gira estas páginas (a partir de 1)."""
    doc = pymupdf.open()
    for n in range(1, paginas + 1):
        folha = doc.new_page(width=595, height=842)
        folha.insert_text((72, 100), f"{titulo} — Página {n}", fontsize=28, fontname="helv")
        folha.insert_text((72, 140), "Ação, atenção e conferência de impostos.", fontsize=14, fontname="helv")
        if com_imagem:
            imagem = Image.effect_noise((1600, 1200), 60).convert("RGB")  # ruído: imagem pesada, comprime bem
            from io import BytesIO
            dados = BytesIO()
            imagem.save(dados, "PNG")
            folha.insert_image(pymupdf.Rect(72, 180, 523, 520), stream=dados.getvalue())
        if n in deitada_em:
            folha.set_rotation(90)
    doc.save(destino)
    doc.close()
    return Path(destino)


TABELA = [["Mês", "Receita", "Despesa", "Margem"], ["Janeiro", "1.200,50", "800,00", "33,4%"],
          ["Fevereiro", "1.500,00", "(900,25)", "40,0%"], ["Março", "R$ 1.700,75", "950,10", "44,1%"]]


def pdf_com_tabela(destino):
    """Uma tabela com linhas desenhadas (como a de um balancete exportado), com números no formato brasileiro."""
    doc = pymupdf.open()
    folha = doc.new_page(width=595, height=842)
    folha.insert_text((60, 70), "Resumo do trimestre", fontsize=18, fontname="helv")
    x0, y0, largura, altura = 60, 100, 115, 26
    for i, linha in enumerate(TABELA):
        for j, valor in enumerate(linha):
            celula = pymupdf.Rect(x0 + j * largura, y0 + i * altura, x0 + (j + 1) * largura, y0 + (i + 1) * altura)
            folha.draw_rect(celula, color=(0, 0, 0), width=0.8)
            folha.insert_text((celula.x0 + 6, celula.y1 - 8), valor, fontsize=11, fontname="helv")
    doc.save(destino)
    doc.close()
    return Path(destino)


def imagem(destino, tamanho=(800, 600), cor=(29, 107, 95), texto="Foto"):
    figura = Image.new("RGB", tamanho, cor)
    desenho = ImageDraw.Draw(figura)
    desenho.rectangle([40, 40, tamanho[0] - 40, tamanho[1] - 40], outline="white", width=8)
    desenho.text((80, 80), texto, fill="white")
    figura.save(destino)
    return Path(destino)


def texto(destino):
    Path(destino).write_text("Recibo de entrega de documentos\n\nCliente: Padaria São João\nItens:\n"
                             "- Notas fiscais de março\n- Extratos bancários\n\n" + "Linha de texto. " * 400,
                             encoding="utf-8")
    return Path(destino)


def office(pasta):
    """Word, Excel e PowerPoint pequenos (python-docx, openpyxl, python-pptx)."""
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    documento = Document()
    documento.add_heading("Contrato de prestação de serviços", 1)
    documento.add_paragraph("Cláusula 1: o escritório cuidará da contabilidade mensal.")
    documento.save(Path(pasta) / "contrato.docx")

    planilha = Workbook()
    aba = planilha.active
    aba.append(["Mês", "Receita", "Despesa"])
    for mes, receita, despesa in (("Janeiro", 1200, 800), ("Fevereiro", 1500, 900), ("Março", 1700, 950)):
        aba.append([mes, receita, despesa])
    planilha.save(Path(pasta) / "balancete.xlsx")

    apresentacao = Presentation()
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[1])
    slide.shapes.title.text = "Reunião mensal"
    slide.placeholders[1].text = "Fechamento de março"
    apresentacao.save(Path(pasta) / "reuniao.pptx")


def gerar(pasta=PASTA_PADRAO):
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    pdf_com_paginas(pasta / "relatorio.pdf", 6)
    pdf_com_paginas(pasta / "anexo.pdf", 2, titulo="Anexo")
    pdf_com_paginas(pasta / "digitalizado-pesado.pdf", 3, titulo="Digitalizado", com_imagem=True)
    pdf_com_paginas(pasta / "paginas-deitadas.pdf", 4, titulo="Deitadas", deitada_em=(2, 4))
    pdf_com_tabela(pasta / "balancete-tabela.pdf")
    imagem(pasta / "foto1.jpg")
    imagem(pasta / "foto2.png", (600, 900), (47, 98, 196), "Retrato")
    texto(pasta / "recibo.txt")
    office(pasta)
    with pymupdf.open(pasta / "anexo.pdf") as doc:
        doc.save(pasta / "com-senha.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1234", owner_pw="1234")
    (pasta / "nao-e-pdf.pdf").write_text("isto não é um PDF", encoding="utf-8")
    return pasta


if __name__ == "__main__":
    print("Exemplos em", gerar(sys.argv[1] if len(sys.argv) > 1 else PASTA_PADRAO))
