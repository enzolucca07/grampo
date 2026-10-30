"""PDF para Excel, PowerPoint, texto e imagens."""

import pymupdf
import pytest
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from grampo import conversao, ferramentas
from grampo.erros import ErroAmigavel


@pytest.fixture
def escaneado(tmp_path):
    """PDF só com imagem, sem texto (como um papel escaneado)."""
    caminho = tmp_path / "escaneado.pdf"
    with pymupdf.open() as doc:
        folha = doc.new_page()
        foto = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), False)
        foto.clear_with(200)
        folha.insert_image(folha.rect, pixmap=foto)
        doc.save(caminho)
    return caminho


@pytest.mark.parametrize("texto, valor, formato", [
    ("1.200,50", 1200.5, "#,##0.00"), ("800,00", 800.0, "#,##0.00"), ("(900,25)", -900.25, "#,##0.00"),
    ("R$ 1.700,75", 1700.75, "#,##0.00"), ("33,4%", 0.334, "0.00%"), ("2026", 2026, None),
    ("1.500", 1500, "#,##0"), ("-12", -12, None), ("0", 0, None),
    ("00123", "00123", None), ("123.456.789-00", "123.456.789-00", None), ("Janeiro", "Janeiro", None),
    ("01/03/2026", "01/03/2026", None),
])
def test_celula_entende_numero_brasileiro(texto, valor, formato):
    convertido, achado = conversao._celula(texto)
    assert convertido == pytest.approx(valor) if isinstance(valor, float) else convertido == valor
    assert achado == formato


def test_pdf_para_excel_tabela(exemplos, saida):
    destino, tabelas = conversao.pdf_para_excel(exemplos / "balancete-tabela.pdf", saida / "b.xlsx")
    assert tabelas == 1
    aba = load_workbook(destino).worksheets[0]
    assert aba.title == "Pág 1"
    assert [c.value for c in aba[1]] == ["Mês", "Receita", "Despesa", "Margem"]
    assert aba["B2"].value == pytest.approx(1200.5) and aba["B2"].number_format == "#,##0.00"
    assert aba["C3"].value == pytest.approx(-900.25)
    assert aba["D2"].value == pytest.approx(0.334)
    assert aba["A1"].font.bold and aba.freeze_panes == "A2"


def test_sem_vazios():
    assert conversao._sem_vazios([["a", None, "1"], [None, None, None], ["b", "", "2"]]) == [["a", "1"], ["b", "2"]]


def test_pdf_para_excel_tabela_sem_grade(tmp_path, saida):
    """Como o PDF que o Excel exporta: colunas só pelo alinhamento, números alinhados à direita, e um título."""
    caminho = tmp_path / "sem-grade.pdf"
    with pymupdf.open() as doc:
        folha = doc.new_page()
        folha.insert_text((56, 50), "Balancete de março de 2026", fontsize=20)
        linhas = [("Conta", "Saldo anterior", "Saldo atual"), ("Caixa geral", "1.200,00", "980,50"),
                  ("Banco do Brasil", "15.300,10", "16.002,00"), ("Fornecedores", "(2.000,00)", "(1.500,00)")]
        for i, (conta, anterior, atual) in enumerate(linhas):
            y = 90 + i * 16
            folha.insert_text((56, y), conta, fontsize=10)
            for direita, valor in ((260, anterior), (360, atual)):  # alinhado à direita
                folha.insert_text((direita - pymupdf.get_text_length(valor, fontsize=10), y), valor, fontsize=10)
        doc.save(caminho)
    destino, tabelas = conversao.pdf_para_excel(caminho, saida / "s.xlsx")
    assert tabelas == 1
    valores = list(load_workbook(destino).worksheets[0].iter_rows(values_only=True))
    assert valores[0] == ("Balancete de março de 2026", None, None)
    assert valores[1] == ("Conta", "Saldo anterior", "Saldo atual")
    assert valores[3][0] == "Banco do Brasil" and valores[3][1:] == pytest.approx((15300.10, 16002.0))
    assert valores[4][1:] == pytest.approx((-2000.0, -1500.0))


def test_pdf_para_excel_sem_tabela_vira_texto(exemplos, saida):
    destino, tabelas = conversao.pdf_para_excel(exemplos / "anexo.pdf", saida / "a.xlsx")
    assert tabelas == 0
    planilha = load_workbook(destino)
    assert [a.title for a in planilha.worksheets] == ["Pág 1 (texto)", "Pág 2 (texto)"]
    assert "Página 1" in planilha.worksheets[0]["A1"].value


def test_pdf_escaneado_explica(escaneado, saida):
    with pytest.raises(ErroAmigavel, match="escaneado"):
        conversao.pdf_para_excel(escaneado, saida / "x.xlsx")
    with pytest.raises(ErroAmigavel, match="escaneado"):
        conversao.pdf_para_texto(escaneado, saida / "x.txt")


def test_pdf_para_powerpoint(exemplos, saida):
    destino = conversao.pdf_para_powerpoint(exemplos / "paginas-deitadas.pdf", saida / "p.pptx", dpi=50)
    apresentacao = Presentation(destino)
    assert len(apresentacao.slides) == 4
    assert apresentacao.slide_height > apresentacao.slide_width  # primeira página em pé
    for slide in apresentacao.slides:
        foto = slide.shapes[0]
        assert foto.shape_type == 13  # imagem
        assert foto.left >= 0 and foto.top >= 0
        assert foto.left + foto.width <= apresentacao.slide_width + 1
        assert foto.top + foto.height <= apresentacao.slide_height + 1


def test_pdf_para_texto(exemplos, saida):
    destino = conversao.pdf_para_texto(exemplos / "relatorio.pdf", saida / "r.txt")
    texto = destino.read_text(encoding="utf-8-sig")
    assert "--- Página 6 ---" in texto and "Ação, atenção" in texto
    assert destino.read_bytes().startswith(b"\xef\xbb\xbf")


def test_extrair_imagens(exemplos, saida):
    imagens = conversao.extrair_imagens(exemplos / "digitalizado-pesado.pdf", saida, "dig")
    assert len(imagens) == 3
    assert imagens[0].name.startswith("dig - imagem 01 (página 1).")
    with Image.open(imagens[0]) as imagem:
        assert imagem.size == (1600, 1200)


def test_extrair_imagens_sem_imagens(exemplos, saida):
    with pytest.raises(ErroAmigavel, match="não tem fotos"):
        conversao.extrair_imagens(exemplos / "anexo.pdf", saida, "a")


def test_pdf_para_png(exemplos, saida):
    imagens = conversao.pdf_para_jpg(exemplos / "anexo.pdf", saida, "anexo", dpi=50, formato="png")
    assert [i.suffix for i in imagens] == [".png", ".png"]
    with Image.open(imagens[0]) as imagem:
        assert imagem.format == "PNG"


def test_ferramentas_novas_pelo_catalogo(exemplos, saida):
    def nomes(resultado):
        return [p.replace("\\", "/").split("/")[-1] for p in resultado["arquivos"]]

    tabela = [(exemplos / "balancete-tabela.pdf", "balancete.pdf")]
    excel = ferramentas.executar("pdf_excel", tabela, {}, saida)
    assert nomes(excel) == ["balancete (convertido).xlsx"] and "1 tabela encontrada" in excel["mensagem"]
    assert nomes(ferramentas.executar("pdf_powerpoint", tabela, {}, saida)) == ["balancete (convertido).pptx"]
    assert nomes(ferramentas.executar("pdf_texto", tabela, {}, saida)) == ["balancete (texto).txt"]
    assert nomes(ferramentas.executar("pdf_jpg", tabela, {"formato": "png", "dpi": "100"}, saida)) == [
        "balancete (imagem).png"]
    extraidas = ferramentas.executar("extrair_imagens", [(exemplos / "digitalizado-pesado.pdf", "dig.pdf")], {},
                                     saida)
    assert len(extraidas["arquivos"]) == 3 and all("dig (imagens extraídas)" in p for p in extraidas["arquivos"])
